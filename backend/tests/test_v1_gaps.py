from datetime import date
from decimal import Decimal
from io import BytesIO
from unittest.mock import MagicMock, patch

from openpyxl import load_workbook
from sqlalchemy import select

from app.attachments.minio_client import StoredObject
from app.db.models.expense import ExpenseClaim, ExpenseEvent, ExpensePayment
from scripts.reconcile_payments import audit, repair
from tests.conftest import login_as, make_user


def make_claim(client, *, amount="100", currency="CNY", status="APPROVED", category="CDN"):
    created = client.post("/ext/api/v1/admin/expenses", json={
        "expense_date": date.today().isoformat(), "category": category,
        "currency": currency, "pay_type": "COMPANY_DIRECT",
        "items": [{"description": "service", "quantity": 1, "unit_price": amount}],
    })
    assert created.status_code == 201, created.text
    claim_id = created.json()["id"]
    if status != "DRAFT":
        assert client.post(f"/ext/api/v1/admin/expenses/{claim_id}/submit").status_code == 200
    if status == "APPROVED":
        assert client.post(f"/ext/api/v1/admin/expenses/{claim_id}/approve").status_code == 200
    elif status == "REJECTED":
        assert client.post(f"/ext/api/v1/admin/expenses/{claim_id}/reject", json={"reason": "no"}).status_code == 200
    return claim_id


def test_partial_payment_balances_and_rejected_writes(client, db):
    admin = make_user(db, sub2api_id=300, role="admin")
    login_as(client, db, admin)
    claim_id = make_claim(client)
    url = f"/ext/api/v1/admin/expenses/{claim_id}/payments"
    assert client.post(url, json={"amount": "30", "mark_paid": True}).status_code == 400
    assert client.post(url, json={"amount": "0"}).status_code == 422
    assert client.post(url, json={"amount": "-1"}).status_code == 422
    assert client.post(url, json={"amount": "1.001"}).status_code == 422
    assert client.post(url, json={"amount": "30", "currency": "USD"}).status_code == 400
    assert client.get(f"/ext/api/v1/admin/expenses/{claim_id}").json()["payments"] == []

    partial = client.post(url, json={"amount": "30"})
    assert partial.status_code == 201, partial.text
    assert partial.json()["status"] == "APPROVED"
    assert Decimal(str(partial.json()["paid_total"])) == 30
    assert Decimal(str(partial.json()["remaining_amount"])) == 70
    assert client.post(url, json={"amount": "71"}).status_code == 400
    assert len(client.get(f"/ext/api/v1/admin/expenses/{claim_id}").json()["payments"]) == 1

    full = client.post(url, json={"amount": "70"})
    assert full.status_code == 201, full.text
    assert full.json()["status"] == "PAID"
    assert Decimal(str(full.json()["remaining_amount"])) == 0
    assert client.post(url, json={"amount": "1"}).status_code == 400
    events = db.scalars(select(ExpenseEvent).where(ExpenseEvent.claim_id == claim_id)).all()
    assert [e.event_type for e in events].count("PAYMENT_ADDED") == 2
    assert [e.event_type for e in events].count("PAID") == 1


def test_foreign_currency_inherited_and_invoice_clear(client, db):
    admin = make_user(db, sub2api_id=301, role="admin")
    login_as(client, db, admin)
    claim_id = make_claim(client, currency="USD")
    payment = client.post(f"/ext/api/v1/admin/expenses/{claim_id}/payments", json={"amount": "100"})
    assert payment.status_code == 201
    assert payment.json()["payments"][0]["currency"] == "USD"
    url = f"/ext/api/v1/admin/expenses/{claim_id}/invoice"
    assert client.put(url, json={"invoice_status": "RECEIVED", "invoice_number": "INV-1",
                                 "seller_name": "Vendor", "tax_amount": "5"}).status_code == 200
    cleared = client.put(url, json={"invoice_number": None, "seller_name": None})
    assert cleared.status_code == 200
    detail = client.get(f"/ext/api/v1/admin/expenses/{claim_id}").json()
    assert detail["invoice_number"] is None
    assert detail["seller_name"] is None
    assert detail["invoice_status"] == "RECEIVED"
    assert Decimal(str(detail["tax_amount"])) == 5


def test_edit_state_and_expense_attachment_permissions(client, db):
    admin = make_user(db, sub2api_id=302, role="admin")
    user = make_user(db, sub2api_id=303, role="user")
    login_as(client, db, admin)
    claim_id = make_claim(client, status="DRAFT")
    patch_body = {
        "expense_date": date.today().isoformat(), "category": "OFFICE",
        "currency": "CNY", "pay_type": "COMPANY_DIRECT",
        "items": [{"description": "paper", "quantity": 2, "unit_price": "10"}],
    }
    edited = client.patch(f"/ext/api/v1/admin/expenses/{claim_id}", json=patch_body)
    assert edited.status_code == 200
    assert Decimal(str(edited.json()["amount_tax_included"])) == 20
    assert client.post(f"/ext/api/v1/admin/expenses/{claim_id}/submit").status_code == 200
    assert client.patch(f"/ext/api/v1/admin/expenses/{claim_id}", json=patch_body).status_code == 400

    storage = MagicMock()
    storage.upload.return_value = StoredObject(
        object_key="expenses/1/x/r.pdf", sha256="a" * 64,
        file_size=4, mime_type="application/pdf", file_name="r.pdf",
    )
    storage.presigned_get_public.return_value = None
    storage.get_bytes.return_value = b"%PDF"
    with patch("app.expenses.router.get_storage", return_value=storage), patch(
        "app.attachments.router.get_storage", return_value=storage
    ):
        uploaded = client.post(
            f"/ext/api/v1/admin/expenses/{claim_id}/attachments",
            params={"attachment_type": "INVOICE"},
            files={"file": ("r.pdf", b"%PDF", "application/pdf")},
        )
        assert uploaded.status_code == 201, uploaded.text
        att = client.get(f"/ext/api/v1/admin/expenses/{claim_id}").json()["attachments"][0]
        assert att["attachment_type"] == "INVOICE"
        download = f"/ext/api/v1/attachments/{att['id']}/download?source=expense"
        assert client.get(download).content == b"%PDF"
        client.cookies.clear()
        login_as(client, db, user)
        assert client.get(download).status_code == 403
        assert client.get(f"/ext/api/v1/admin/expenses/{claim_id}").status_code == 403


def test_currency_status_reports_and_exports(client, db):
    admin = make_user(db, sub2api_id=304, role="admin")
    login_as(client, db, admin)
    cny = make_claim(client, amount="100")
    client.post(f"/ext/api/v1/admin/expenses/{cny}/payments", json={"amount": "30"})
    client.put(f"/ext/api/v1/admin/expenses/{cny}/invoice", json={
        "invoice_status": "RECEIVED", "tax_amount": "6", "deductible": True,
    })
    usd = make_claim(client, amount="50", currency="USD")
    client.post(f"/ext/api/v1/admin/expenses/{usd}/payments", json={"amount": "50"})
    make_claim(client, amount="20", status="REJECTED")
    make_claim(client, amount="25", status="SUBMITTED")
    params = {"period": "year", "year": date.today().year}

    groups = {r["currency"]: r for r in client.get("/ext/api/v1/admin/reports/summary", params=params).json()["currencies"]}
    assert groups["CNY"]["total_amount"] == 100
    assert groups["CNY"]["tax_amount"] == 6
    assert groups["CNY"]["by_status"]["REJECTED"]["amount"] == 20
    assert groups["CNY"]["by_status"]["SUBMITTED"]["amount"] == 25
    assert groups["USD"]["total_amount"] == 50
    payment = {r["currency"]: r for r in client.get("/ext/api/v1/admin/reports/payment-status", params=params).json()["currencies"]}
    assert payment["CNY"]["paid_amount"] == 30
    assert payment["CNY"]["unpaid_approved_amount"] == 70
    assert payment["CNY"]["rejected_amount"] == 20
    assert payment["USD"]["paid_amount"] == 50
    invoice = {r["currency"]: r for r in client.get(
        "/ext/api/v1/admin/reports/invoice-tax", params=params).json()["currencies"]}
    assert invoice["CNY"]["with_invoice_amount"] == 100
    assert invoice["CNY"]["deductible_tax"] == 6
    for endpoint in ("by-month", "by-category", "by-supplier", "by-cost-center"):
        rows = client.get(f"/ext/api/v1/admin/reports/{endpoint}", params=params)
        assert rows.status_code == 200
        assert {r["currency"] for r in rows.json()} == {"CNY", "USD"}
    filtered = client.get("/ext/api/v1/admin/reports/summary", params={**params, "currency": "USD"})
    assert [r["currency"] for r in filtered.json()["currencies"]] == ["USD"]
    csv_response = client.get("/ext/api/v1/admin/reports/export",
                              params={**params, "view": "payment", "format": "csv"})
    assert csv_response.status_code == 200
    assert "unpaid_approved_amount" in csv_response.text
    xlsx = client.get("/ext/api/v1/admin/reports/export",
                      params={**params, "view": "all", "format": "xlsx"})
    workbook = load_workbook(BytesIO(xlsx.content), read_only=True)
    assert set(workbook.sheetnames) == {"detail", "summary", "month", "category",
                                         "supplier", "cost_center", "payment", "invoice"}
    assert workbook["summary"].max_row >= 3
    empty = client.get("/ext/api/v1/admin/reports/export", params={
        "period": "year", "year": 2099, "format": "csv", "view": "summary",
    })
    assert empty.status_code == 200
    assert "currency,count,total_amount" in empty.text


def test_legacy_reconciliation_requires_expected_values(client, db):
    admin = make_user(db, sub2api_id=305, role="admin")
    login_as(client, db, admin)
    claim_id = make_claim(client)
    claim = db.get(ExpenseClaim, claim_id)
    claim.status = "PAID"
    db.commit()
    findings = audit(db)
    assert len(findings) == 1
    assert findings[0]["reason"] == "paid_without_full_payment"
    detail_csv = client.get("/ext/api/v1/admin/reports/export", params={
        "period": "year", "year": date.today().year, "format": "csv", "view": "detail",
    })
    assert detail_csv.status_code == 200
    assert "payment_reconciliation_required" in detail_csv.text
    assert "True" in detail_csv.text
    db.refresh(claim)
    assert claim.status == "PAID"  # audit is read-only
    base = {
        "claim_no": claim.claim_no, "expected_amount": "100.00",
        "expected_paid": "0", "action": "restore_approved",
        "reason": "verified no external payment", "reference_no": "",
    }
    from pytest import raises
    with raises(ValueError, match="expected amounts changed"):
        repair(db, [{**base, "expected_paid": "1"}], admin.id)
    db.rollback()
    repair(db, [base], admin.id)
    db.refresh(claim)
    assert claim.status == "APPROVED"
    assert audit(db) == []
    assert any(e.event_type == "PAYMENT_RECONCILED" for e in db.scalars(
        select(ExpenseEvent).where(ExpenseEvent.claim_id == claim_id)))

    other_id = make_claim(client)
    other = db.get(ExpenseClaim, other_id)
    other.status = "PAID"
    db.commit()
    repair(db, [{
        "claim_no": other.claim_no, "expected_amount": "100", "expected_paid": "0",
        "action": "record_payment", "reason": "verified external transfer",
        "reference_no": "legacy-bank-ref", "payment_date": "2026-09-22T09:00:00+08:00",
    }], admin.id)
    db.refresh(other)
    assert other.status == "PAID"
    assert len(db.scalars(select(ExpensePayment).where(ExpensePayment.claim_id == other_id)).all()) == 1
    assert audit(db) == []
