import uuid
from datetime import date, datetime, timezone, timedelta
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

import httpx
import pytest
from openpyxl import load_workbook
from sqlalchemy import select, func

from app.db.models.expense import ExpenseClaim, ExpenseAttachment, ExpensePayment
from app.db.models.ledger import ExpenseRecord, RateQuote, Reimbursement, LedgerAttachment, LedgerEvent
from app.ledger import rates
from tests.conftest import login_as, make_user

ROOT = '/ext/api/v1/admin/expense-records'
BATCH = '/ext/api/v1/admin/reimbursements'


@pytest.fixture()
def admin(client, db):
    actor = make_user(db, sub2api_id=8901, role='admin', username='录入管理员')
    login_as(client, db, actor)
    return actor


def payload(**changes):
    return dict(title='服务器年付', expense_date='2026-10-01', category='CLOUD_SERVER', payer=' 张三 ',
                amount='100.00', currency='CNY', idempotency_key=str(uuid.uuid4()), **changes)


def create(client, **changes):
    body = payload()
    body.update(changes)
    response = client.post(ROOT, json=body)
    assert response.status_code == 201, response.text
    return response.json()


def quote(client, monkeypatch, rate='7.10'):
    monkeypatch.setattr(rates.provider, 'latest', lambda: (Decimal(rate), date(2026, 10, 9)))
    result = client.get(ROOT + '/rate', params={'refresh': True})
    assert result.status_code == 200, result.text
    return result.json()


def batch_body(records, **changes):
    body = dict(records=[{'id': r['id'], 'version': r['version']} for r in records], actual_amount='100.00',
                reimbursement_date='2026-11-02', idempotency_key=str(uuid.uuid4()))
    body.update(changes)
    return body


def test_cny_default_immediate_month_and_idempotency(client, db, admin):
    body = payload()
    a = client.post(ROOT, json=body).json()
    b = client.post(ROOT, json=body).json()
    assert a['id'] == b['id'] and a['payer'] == '张三' and a['created_by'] == admin.id
    assert a['status'] == 'PENDING' and a['rmb_amount'] == '100.00'
    monthly = client.get(ROOT + '/monthly', params={'month':'2026-10'}).json()
    assert monthly['total_amount'] == '100.00' and monthly['people'][0]['pending'] == '100.00'
    assert client.get(ROOT, params={'month':'2026-11'}).json()['total'] == 0
    body['amount'] = '101.00'
    assert client.post(ROOT, json=body).status_code == 409
    assert db.scalar(select(func.count()).select_from(ExpenseRecord)) == 1


@pytest.mark.parametrize('field,value', [('amount','0'),('amount','-1'),('amount','1.001'),('amount','NaN'),('amount','Infinity'),('amount','1e3'),('amount',1),('amount','10000000000000000'),('payer','  '),('title',''),('currency','HKD'),('category','STAFF_ADVANCE')])
def test_invalid_values_rejected(client, admin, field, value):
    body = payload(); body[field] = value
    assert client.post(ROOT, json=body).status_code == 422


def test_usd_snapshot_manual_edit_and_all_outputs(client, db, admin, monkeypatch):
    q = quote(client, monkeypatch)
    automatic = create(client, currency='USD', amount='20.00', quote_id=q['id'])
    assert automatic['rmb_amount'] == '142.00'
    manual = create(client, currency='USD', amount='20.00', amount_mode='MANUAL', rmb_amount='143.26', quote_id=q['id'])
    quote(client, monkeypatch, '8.00')
    assert client.get(f"{ROOT}/{automatic['id']}").json()['rmb_amount'] == '142.00'
    unchanged = payload(); unchanged.pop('idempotency_key'); unchanged.update(currency='USD', amount='20.00', amount_mode='MANUAL', rmb_amount='143.26', quote_id=q['id'], version=manual['version'])
    edited = client.patch(f"{ROOT}/{manual['id']}", json=unchanged)
    assert edited.status_code == 200, edited.text
    assert Decimal(edited.json()['reference_rate']) == Decimal('7.10')
    report = client.get(ROOT + '/monthly', params={'month':'2026-10'}).json()
    assert report['total_amount'] == '285.26'
    reimb = client.post(BATCH, json=batch_body([edited.json()], actual_amount='143.26'))
    assert reimb.status_code == 201, reimb.text
    assert reimb.json()['items'][0]['expected_amount'] == '143.26'
    assert client.get(ROOT + '/monthly', params={'month':'2026-11'}).json()['total_amount'] == '0.00'
    book = load_workbook(BytesIO(client.get(ROOT + '/export', params={'month':'2026-10'}).content))
    assert book.sheetnames == ['支出明细','分类汇总','人员汇总']
    sheet = book['支出明细']
    values = list(sheet.iter_rows(min_row=6, values_only=True))
    assert sum(Decimal(str(v[13])) for v in values) == Decimal('285.26')
    assert any(v[13] == 143.26 and v[15] == 143.26 for v in values)


def test_failure_manual_without_fake_rate_and_cny_switch(client, admin, monkeypatch):
    def failure(): raise httpx.ConnectError('offline')
    monkeypatch.setattr(rates.provider, 'latest', failure)
    assert client.get(ROOT + '/rate', params={'refresh': True}).status_code == 503
    row = create(client, currency='USD', amount='20', amount_mode='MANUAL', rmb_amount='143.26')
    assert row['reference_rate'] is None and row['rate_source'] == '人工填写'
    body = payload(); body.pop('idempotency_key'); body.update(version=row['version'], amount='20', rmb_amount='999', quote_id='stale', currency='CNY')
    result = client.patch(f"{ROOT}/{row['id']}", json=body)
    assert result.status_code == 200, result.text
    assert result.json()['rmb_amount'] == '20.00' and result.json()['quote_id'] is None and result.json()['rate_source'] is None
    assert client.post(ROOT, json={**payload(), 'currency':'USD'}).status_code == 400


def test_half_cent_rounding_and_decimal_sum(client, admin, monkeypatch):
    q = quote(client, monkeypatch, '7.105')
    assert create(client, currency='USD', amount='1.00', quote_id=q['id'])['rmb_amount'] == '7.11'
    for _ in range(3): create(client, amount='0.10')
    assert client.get(ROOT+'/monthly', params={'month':'2026-10'}).json()['total_amount'] == '7.41'
    q = quote(client, monkeypatch, '0.0001')
    assert client.post(ROOT, json={**payload(), 'amount':'0.01', 'currency':'USD','quote_id':q['id']}).status_code == 400


def test_batch_allocation_revoke_and_optimistic_edit(client, db, admin):
    a, b = create(client), create(client, currency='USD', amount='20', rmb_amount='200', amount_mode='MANUAL')
    body = batch_body([b,a], actual_amount='299.00', difference_reason='手续费调整')
    preview = client.post(BATCH + '/preview', json=body)
    assert preview.status_code == 200, preview.text
    assert [r['actual_amount'] for r in preview.json()] == ['99.67','199.33']
    response = client.post(BATCH, json=body)
    assert response.status_code == 201, response.text
    batch = response.json()
    assert client.post(BATCH, json=body).json()['id'] == batch['id']
    assert sum(Decimal(i['difference']) for i in batch['items']) == Decimal('-1.00')
    assert client.get(ROOT+'/monthly', params={'month':'2026-10'}).json()['total_amount'] == '300.00'
    a = client.get(f"{ROOT}/{a['id']}").json()
    assert a['status'] == 'REIMBURSED'
    assert client.post(f"{ROOT}/{a['id']}/void", json={'version':a['version'],'reason':'更正'}).status_code == 409
    edit = payload(); edit.pop('idempotency_key'); edit.update(version=a['version'], amount='200')
    assert client.patch(f"{ROOT}/{a['id']}", json=edit).status_code == 409
    revoked = client.post(f"{BATCH}/{batch['id']}/revoke", json={'version':1,'reason':'登记错误'})
    assert revoked.status_code == 200, revoked.text
    assert client.get(f"{ROOT}/{a['id']}").json()['status'] == 'PENDING'
    assert client.get(f"{ROOT}/{b['id']}").json()['status'] == 'PENDING'
    assert db.scalar(select(func.count()).select_from(Reimbursement)) == 1
    assert client.patch(f"{ROOT}/{a['id']}", json=edit).json()['code'] == 'VERSION_CONFLICT'
    a = client.get(f"{ROOT}/{a['id']}").json()
    result = client.post(f"{ROOT}/{a['id']}/void", json={'version':a['version'],'reason':'重复记账'})
    assert result.status_code == 200
    assert client.get(ROOT+'/monthly', params={'month':'2026-10'}).json()['total_amount'] == '200.00'
    assert client.get(ROOT, params={'month':'2026-10','voided':True}).json()['total'] == 1


@pytest.mark.parametrize('scenario', ['different_payer','duplicate','not_required','voided','version','difference','zero'])
def test_invalid_batch_no_partial_success(client, db, admin, scenario):
    a = create(client); b = create(client, payer='小张' if scenario=='different_payer' else '张三', needs_reimbursement=scenario!='not_required')
    if scenario=='voided':
        client.post(f"{ROOT}/{b['id']}/void", json={'version':b['version'],'reason':'重复'})
        b = client.get(f"{ROOT}/{b['id']}").json()
    body = batch_body([a,b], actual_amount='200')
    if scenario=='duplicate': body['records'][1] = body['records'][0]
    if scenario=='version': body['records'][1]['version'] = 999
    if scenario=='difference': body['actual_amount']='199'
    if scenario=='zero': body['actual_amount']='0'
    assert client.post(BATCH, json=body).status_code in (400,409,422)
    assert db.scalar(select(func.count()).select_from(Reimbursement)) == 0
    assert client.get(f"{ROOT}/{a['id']}").json()['status'] == 'PENDING'


def test_not_required_pagination_export_formula_and_keywords(client, admin):
    create(client, needs_reimbursement=False, payer='=SUM(1,2)', title='=1+2', merchant='平台', notes='备注关键词')
    for i in range(54): create(client, title=f'订阅{i}', amount='0.10')
    assert client.get(ROOT, params={'month':'2026-10','page':3}).json()['total'] == 55
    assert len(client.get(ROOT, params={'month':'2026-10','page':3}).json()['items']) == 15
    for word in ['=SUM','平台','备注关键词','=1+2']:
        assert client.get(ROOT, params={'month':'2026-10','keyword':word}).json()['total'] == 1
    report = client.get(ROOT+'/monthly', params={'month':'2026-10'}).json()
    person = next(p for p in report['people'] if p['payer']=='=SUM(1,2)')
    assert person['pending']=='0.00' and person['not_required']=='100.00'
    sheet = load_workbook(BytesIO(client.get(ROOT+'/export', params={'month':'2026-10'}).content))['支出明细']
    assert sheet.max_row == 60
    row = next(row for row in sheet.iter_rows(min_row=6) if row[2].value=='=1+2')
    assert row[2].data_type == 's' and row[5].data_type == 's' and row[13].data_type == 'n'
    assert row[13].number_format=='0.00' and row[15].value is None


def old_claim(db, admin, **changes):
    values = dict(claim_no='LEGACY-'+uuid.uuid4().hex[:10], applicant_user_id=admin.id, expense_date=date(2026,9,3),
                  category='CDN', currency='USD', amount_tax_included=Decimal('20'), pay_type='PERSONAL_ADVANCE', status='PAID', invoice_status='NONE')
    values.update(changes)
    row = ExpenseClaim(**values); db.add(row); db.commit(); return row


def test_history_explicit_confirmation_unique_and_original_preserved(client, db, admin):
    old = old_claim(db, admin)
    assert client.get(ROOT+'/monthly', params={'month':''}).json()['unconfirmed_count']==1
    assert client.get(ROOT+'/monthly', params={'month':''}).json()['total_amount']=='0.00'
    body = {**payload(), 'expense_date':'2026-09-03','currency':'USD','amount':'20','amount_mode':'MANUAL','rmb_amount':'140.00','history_settled':False}
    target = f'{ROOT}/history/{old.id}/confirm'
    created = client.post(target,json=body)
    assert created.status_code==201, created.text
    assert created.json()['status']=='PENDING' and created.json()['reference_rate'] is None
    assert client.post(target,json=body).json()['id']==created.json()['id']
    body['payer']='李四'
    assert client.post(target,json=body).status_code==409
    db.refresh(old)
    assert old.currency=='USD' and old.status=='PAID' and old.amount_tax_included==Decimal('20')
    assert client.get(ROOT+'/history').json()['unconfirmed_count']==0
    assert client.get('/ext/api/v1/admin/expenses/'+str(old.id)).status_code==200


def test_history_anomalies_and_paid_evidence(client, db, admin):
    old=old_claim(db,admin,currency='HKD',amount_tax_included=0,status='CANCELLED')
    db.add(ExpensePayment(claim_id=old.id,amount=Decimal('10'),currency='USD',paid_at=datetime.now(timezone.utc),created_by_user_id=admin.id)); db.commit()
    report=client.get(ROOT+'/history').json()
    assert set(report['items'][0]['anomalies']) >= {'其他历史币种','零或负金额','原单已取消','跨币种或非正付款'}
    target=f'{ROOT}/history/{old.id}/confirm'
    body={**payload(),'history_settled':True}
    assert client.post(target,json=body).status_code==400
    body.update(actual_reimbursement='99',reimbursement_date='2026-10-03',difference_reason='核对转账记录')
    assert client.post(target,json=body).json()['status']=='REIMBURSED'


def test_legacy_writes_retired_and_permissions(client,db,admin):
    assert client.post('/ext/api/v1/admin/expenses',json={}).status_code==409
    old=old_claim(db,admin)
    for suffix in ('submit','approve','cancel','payments','attachments'):
        assert client.post(f'/ext/api/v1/admin/expenses/{old.id}/{suffix}',json={}).status_code==409
    assert client.get('/ext/api/v1/admin/reports/summary',params={'period':'month','year':2026,'month':10}).headers['x-report-scope']=='legacy-claims-original-currency'
    user=make_user(db,sub2api_id=8902,role='user'); login_as(client,db,user)
    for suffix in ('','/meta','/history','/monthly','/rate','/export','/1'):
        assert client.get(ROOT+suffix).status_code==403
    assert client.post(ROOT,json=payload()).status_code==403
    assert client.post(BATCH,json=batch_body([])).status_code==403
    for source in ('expense','expense-record','reimbursement'):
        assert client.get('/ext/api/v1/attachments/1/download',params={'source':source}).status_code==403


def test_quote_cache_refresh_failure_and_provider_validation(client,db,admin,monkeypatch):
    q=quote(client,monkeypatch)
    assert client.get(ROOT+'/rate').json()['id']==q['id']
    def offline(): raise httpx.ConnectError('offline')
    monkeypatch.setattr(rates.provider,'latest',offline)
    assert client.get(ROOT+'/rate',params={'refresh':True}).status_code==503
    assert client.get(ROOT+'/rate').json()['fetched_at']==q['fetched_at']
    cached=db.get(RateQuote,q['id']); cached.fetched_at=datetime.now(timezone.utc)-timedelta(minutes=31);db.commit()
    assert client.get(ROOT+'/rate').status_code==503


@pytest.mark.parametrize('data', [dict(base='EUR',quote='CNY',rate=7,date='2026-10-09'),dict(base='USD',quote='CNY',rate=-1,date='2026-10-09'),dict(base='USD',quote='CNY',rate='NaN',date='2026-10-09')])
def test_provider_rejects_invalid_response(data):
    with patch('app.ledger.rates.httpx.Client') as factory:
        factory.return_value.__enter__.return_value.get.return_value.text=__import__('json').dumps(data)
        with pytest.raises(ValueError): rates.FrankfurterProvider().latest()


def test_ledger_and_batch_attachment_sources_and_legacy_link(client, db, admin):
    from app.attachments.minio_client import StoredObject
    row=create(client)
    stored=StoredObject('ledger/private/proof','f'*64,8,'text/plain','proof.txt')
    with patch('app.ledger.router.get_storage') as storage:
        storage.return_value.upload.return_value=stored
        result=client.post(f"{ROOT}/{row['id']}/attachments",files={'file':('proof.txt',b'evidence','text/plain')})
        assert result.status_code==201,result.text
        expense_attachment_id=result.json()['id']
        batch=client.post(BATCH,json=batch_body([row])).json()
        result=client.post(f"{BATCH}/{batch['id']}/attachments",files={'file':('proof.txt',b'evidence','text/plain')})
        assert result.status_code==201,result.text
        batch_attachment_id=result.json()['id']
    with patch('app.attachments.router.get_storage') as storage:
        storage.return_value.get_bytes.return_value=b'evidence'
        for attachment_id,source in [(expense_attachment_id,'expense-record'),(batch_attachment_id,'reimbursement')]:
            url=f'/ext/api/v1/attachments/{attachment_id}/download'
            assert client.get(url,params={'source':source,'proxy':True}).content==b'evidence'
            wrong='reimbursement' if source=='expense-record' else 'expense-record'
            assert client.get(url,params={'source':wrong,'proxy':True}).status_code==404
        assert client.get('/ext/api/v1/attachments/1/download',params={'source':'unknown'}).status_code==400
    old=old_claim(db,admin,currency='CNY')
    legacy_attachment=ExpenseAttachment(claim_id=old.id,uploader_user_id=admin.id,file_name='old.pdf',object_key='unchanged/key',mime_type='application/pdf',file_size=8,sha256='e'*64)
    db.add(legacy_attachment);db.commit()
    confirmed=client.post(f'{ROOT}/history/{old.id}/confirm',json={**payload(),'history_settled':False}).json()
    assert confirmed['attachments'][0]['source']=='expense' and confirmed['attachments'][0]['id']==legacy_attachment.id
    with patch('app.attachments.router.get_storage') as storage:
        storage.return_value.get_bytes.return_value=b'%PDF-old'
        assert client.get(f'/ext/api/v1/attachments/{legacy_attachment.id}/download',params={'source':'expense','proxy':True}).content==b'%PDF-old'
        storage.return_value.get_bytes.assert_called_with('unchanged/key')


def test_upload_failure_preserves_record_and_retry_only_uploads(client,db,admin):
    row=create(client)
    with patch('app.ledger.router.get_storage') as storage:
        storage.return_value.upload.side_effect=RuntimeError('offline')
        response=client.post(f"{ROOT}/{row['id']}/attachments",files={'file':('proof.txt',b'receipt','text/plain')})
        assert response.status_code==503
    assert db.scalar(select(func.count()).select_from(ExpenseRecord))==1
    assert client.get(f"{ROOT}/{row['id']}").json()['rmb_amount']=='100.00'


def test_reimbursed_core_field_locks_and_noncore_audit(client,admin):
    row=create(client);client.post(BATCH,json=batch_body([row]))
    row=client.get(f"{ROOT}/{row['id']}").json()
    body=payload();body.pop('idempotency_key');body['version']=row['version']
    for field,value in [('amount','200'),('payer','小张'),('expense_date','2026-09-01'),('needs_reimbursement',False),('currency','USD')]:
        change={**body,field:value}
        if field=='currency':change.update(amount_mode='MANUAL',rmb_amount='100')
        assert client.patch(f"{ROOT}/{row['id']}",json=change).status_code==409
    body['notes']='补充用途'
    edited=client.patch(f"{ROOT}/{row['id']}",json=body)
    assert edited.status_code==200,edited.text
    update=edited.json()['events'][-1]
    assert update['old_value']['notes']=='' and update['new_value']['notes']=='补充用途' and update['actor_name']=='录入管理员'


def test_my_records_filter_uses_creator_not_free_text_payer(client,db,admin):
    create(client,payer='李四')
    other=make_user(db,sub2api_id=8903,role='admin',username='另一管理员');login_as(client,db,other)
    create(client,payer='录入管理员')
    mine=client.get(ROOT,params={'month':'2026-10','mine':True}).json()
    assert mine['total']==1 and mine['items'][0]['payer']=='录入管理员' and mine['items'][0]['created_by']==other.id


def test_provider_parses_decimal_json_exactly_and_checks_direction():
    with patch('app.ledger.rates.httpx.Client') as factory:
        factory.return_value.__enter__.return_value.get.return_value.text='{"date":"2026-10-09","base":"USD","quote":"CNY","rate":7.105123456789}'
        value,effective=rates.FrankfurterProvider().latest()
        assert value==Decimal('7.105123456789') and effective==date(2026,10,9)
        factory.assert_called_once_with(timeout=5.0,follow_redirects=False)


def test_auto_server_amount_and_nonmonetary_edit_preserve_saved_snapshot(client,db,admin,monkeypatch):
    q=quote(client,monkeypatch)
    row=create(client,currency='USD',amount='20',rmb_amount='999',quote_id=q['id'])
    assert row['rmb_amount']=='142.00'
    # Even a changed provider quote row cannot silently revalue an existing entry.
    cached=db.get(RateQuote,q['id']);cached.rate=Decimal('8');db.commit()
    body=payload();body.pop('idempotency_key');body.update(currency='USD',amount='20',quote_id=q['id'],version=row['version'],notes='补记用途')
    response=client.patch(f"{ROOT}/{row['id']}",json=body)
    assert response.status_code==200,response.text
    assert response.json()['rmb_amount']=='142.00' and Decimal(response.json()['reference_rate'])==Decimal('7.10')
