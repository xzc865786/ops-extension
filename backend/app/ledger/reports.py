from decimal import Decimal
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font
from app.ledger import service


def monthly(db, filters, actor):
    rows = db.execute(service.query_records(filters, actor)).all()
    items = [service.record_output(r, i, b) for r, i, b in rows]
    categories, people = {}, {}
    total = Decimal(0)
    for r, i, b in rows:
        total += r.rmb_amount
        cat = categories.setdefault(r.category, {"category": r.category, "label": service.EXPENSE_CATEGORY_LABELS.get(r.category, r.category), "amount": Decimal(0), "count": 0})
        cat["amount"] += r.rmb_amount
        cat["count"] += 1
        person = people.setdefault(r.payer, {"payer": r.payer, **{k: Decimal(0) for k in ("required", "pending", "settled", "actual", "difference", "not_required")}})
        if not r.needs_reimbursement:
            person["not_required"] += r.rmb_amount
        else:
            person["required"] += r.rmb_amount
            if i:
                person["settled"] += r.rmb_amount
                person["actual"] += i.actual_amount
                person["difference"] += i.difference
            else:
                person["pending"] += r.rmb_amount
    def output(values):
        return [{k: service.dec(v) if isinstance(v, Decimal) else v for k, v in row.items()} for row in values]
    pending = sum((r.rmb_amount for r, i, _ in rows if r.needs_reimbursement and not i), Decimal(0))
    previous = service.previous_range(filters)
    previous_total = None
    if previous:
        scope = {**filters, "month": "", "date_from": previous[0].isoformat(), "date_to": previous[1].isoformat()}
        previous_total = service.dec(sum((r.rmb_amount for r, _, _ in db.execute(service.query_records(scope, actor))), Decimal(0)))
    return {"total_amount": service.dec(total), "count": len(items), "items": items,
            "pending_amount": service.dec(pending), "previous_total_amount": previous_total,
            "previous_range": [d.isoformat() for d in previous] if previous else None,
            "categories": output(categories.values()), "people": output(people.values()),
            "filters": filters, "unconfirmed_count": service.unconfirmed_count(db),
            "scope": "按支出日期筛选；实际已报销额为这些支出的有效报销分配额，不是所选期间的转账流水"}


STATUS_LABELS = {"PENDING": "待报销", "REIMBURSED": "已报销", "NOT_REQUIRED": "无需报销", "VOIDED": "作废"}
DETAIL_COLUMNS = [
    ("id", "记录编号"), ("expense_date", "支出日期"), ("title", "购买内容"), ("category_label", "分类"),
    ("merchant", "商家"), ("payer", "付款人"), ("currency", "币种"), ("amount", "原币金额"),
    ("reference_rate", "参考汇率（1美元兑人民币）"), ("rate_date", "汇率有效日期"), ("rate_source", "参考汇率来源"),
    ("rate_fetched_at", "汇率获取时间"), ("amount_mode", "金额模式"), ("rmb_amount", "人民币记账额"),
    ("status", "报销状态"), ("actual_reimbursement", "实际报销额"), ("difference", "差额"),
    ("reimbursement_date", "报销日期"), ("batch_id", "报销批次号"), ("notes", "备注"),
]


def excel(report):
    book = Workbook()
    book.remove(book.active)
    sheets = [
        ("支出明细", DETAIL_COLUMNS, report["items"]),
        ("分类汇总", [("label", "分类"), ("amount", "人民币金额"), ("count", "笔数")], report["categories"]),
        ("人员汇总", [("payer", "付款人"), ("required", "需报销记账额"), ("pending", "当前待报销额"),
                      ("settled", "已结清记账额"), ("actual", "实际已报销额"), ("difference", "差额"), ("not_required", "无需报销金额")], report["people"]),
    ]
    money_fields = {"amount", "rmb_amount", "actual_reimbursement", "difference", "required", "pending", "settled", "actual", "not_required"}
    for name, columns, rows in sheets:
        sheet = book.create_sheet(name)
        notes = ["金额单位：人民币；" + report["scope"], "筛选条件：" + str(report["filters"])]
        for n, note in enumerate(notes, 1):
            cell = sheet.cell(n, 1, note)
            cell.data_type = "s"
        for n, (_, label) in enumerate(columns, 1):
            cell = sheet.cell(5, n, label)
            cell.font = Font(bold=True)
            sheet.column_dimensions[cell.column_letter].width = 22
        for rn, row in enumerate(rows, 6):
            for cn, (key, _) in enumerate(columns, 1):
                value = row.get(key)
                if key == "status":
                    value = STATUS_LABELS.get(value, value)
                elif key == "amount_mode":
                    value = "已手动调整" if value == "MANUAL" else "自动"
                cell = sheet.cell(rn, cn)
                if value is None:
                    continue
                if key in money_fields or key == "reference_rate":
                    cell.value = Decimal(str(value))
                    cell.number_format = "0.00000000" if key == "reference_rate" else "0.00"
                elif key == "count":
                    cell.value = value
                else:
                    cell.value = str(value)
                    cell.data_type = "s"
        sheet.freeze_panes = "A6"
        sheet.auto_filter.ref = f"A5:{sheet.cell(max(5, sheet.max_row), len(columns)).coordinate}"
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()
